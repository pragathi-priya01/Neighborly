import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from report_helpers import (
    set_cell_margins, set_cell_background, set_cell_borders, 
    add_code_box, add_callout_box
)

def add_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    # Configure Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("NEIGHBOURLY: A HYPERLOCAL SERVICE MARKETPLACE")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(100, 116, 139)
    
    # Footer with dynamic page numbering
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Department of Computer Science and Engineering  |  Page ")
    frun.font.name = 'Times New Roman'
    frun.font.size = Pt(9.0)
    frun.font.color.rgb = RGBColor(100, 116, 139)
    
    # Add page number XML field
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)
    
    frun2 = fp.add_run(" of ")
    frun2.font.name = 'Times New Roman'
    frun2.font.size = Pt(9.0)
    frun2.font.color.rgb = RGBColor(100, 116, 139)
    fldNumpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    fp._p.append(fldNumpages)

def format_paragraph(p, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.alignment = align

def add_p(doc, text="", bold=False, italic=False, size=12, color=RGBColor(0,0,0), space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    format_paragraph(p, space_after=space_after, line_spacing=1.15, align=align)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p

def add_heading1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return p

def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_pre = p.add_run(bold_prefix + ": ")
    r_pre.font.name = 'Times New Roman'
    r_pre.font.size = Pt(12)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(15, 23, 42)
    
    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(12)
    r_txt.font.color.rgb = RGBColor(0, 0, 0)

def add_table_styled(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1E293B") # Dark slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        set_cell_borders(hdr_cells[i], "0F172A", "0F172A", "0F172A", "0F172A", sz="8")
        hp = hdr_cells[i].paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrun = hp.runs[0]
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(10.5)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(255, 255, 255)
        
    # Format Body Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=130, right=130)
            set_cell_borders(row_cells[c_idx], "E2E8F0", "E2E8F0", "E2E8F0", "E2E8F0", sz="4")
            cp = row_cells[c_idx].paragraphs[0]
            cp.paragraph_format.line_spacing = 1.05
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or len(str(val)) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            cp.alignment = align
            if len(cp.runs) > 0:
                crun = cp.runs[0]
                crun.font.name = 'Times New Roman'
                crun.font.size = Pt(10.0)
                crun.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def add_figure_image(doc, img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        crun = p_cap.add_run(caption)
        crun.font.name = 'Times New Roman'
        crun.font.size = Pt(10.5)
        crun.font.bold = True
        crun.font.color.rgb = RGBColor(30, 41, 59)
    else:
        print(f"Warning: image {img_path} not found.")

print("Report base layout initialized.")
